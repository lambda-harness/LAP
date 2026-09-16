from pathlib import Path

import hashlib

from invoice_agent import input_files, normalize, workbook
from openpyxl import load_workbook


def test_normalize_and_workbook(tmp_path: Path) -> None:
    row, warnings = normalize({
        "invoice_type": "电子发票", "invoice_number": "12345678901234567890",
        "issue_date": "2026-07-20", "buyer_name": "示例购买方有限公司",
        "buyer_tax_id": "91110000123456789X", "seller_name": "示例销售方有限公司",
        "seller_tax_id": "91310000123456789X", "amount_excluding_tax": 159.25,
        "tax_amount": 9.55, "total_amount": 168.80, "currency": "CNY",
    }, "sample.pdf", 1)
    assert warnings == []
    target = tmp_path / "result.xlsx"
    workbook([row], target)
    book = load_workbook(target, read_only=True, data_only=True)
    values = list(book["发票明细"].iter_rows(values_only=True))
    book.close()
    assert values[1][3] == "12345678901234567890"
    assert values[1][11] == 168.8


def test_invalid_invoice_number_requires_review() -> None:
    row, warnings = normalize({
        "invoice_number": "1234567890123456789", "amount_excluding_tax": 159.25,
        "tax_amount": 9.55, "total_amount": 168.80,
    }, "sample.pdf", 1)
    assert "发票号码长度或格式异常" in warnings
    assert row[13] == "需复核"


def test_multiple_host_artifacts_are_resolved(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    artifacts = []
    for index, suffix in enumerate((".png", ".pdf"), 1):
        opaque = f"artifact-{index}"
        content = f"fixture-{index}".encode()
        (input_dir / opaque).write_bytes(content)
        artifacts.append({
            "name": f"invoice-{index}{suffix}",
            "uri": f"lap://run/input/{opaque}",
            "sha256": hashlib.sha256(content).hexdigest(),
        })

    resolved = input_files({"artifacts": artifacts}, tmp_path)
    assert [name for _, name in resolved] == ["invoice-1.png", "invoice-2.pdf"]
