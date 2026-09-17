"""Extract Chinese VAT invoice details into one LAP-delivered workbook."""

from __future__ import annotations

import base64
import contextlib
import hashlib
import json
import mimetypes
import os
import re
import sys
import urllib.error
import urllib.request
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pymupdf  # type: ignore[import-not-found]
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]

AGENT_ID = "io.github.lambda-harness.vat-invoice-agent"
VERSION = "0.2.1"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_FILES = 20
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 50
INPUT_URI = re.compile(r"^lap://run/input/([A-Za-z0-9._-]+)$")
ALLOWED_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf"}
HEADERS = [
    "来源文件",
    "页码",
    "发票类型",
    "发票号码",
    "开票日期",
    "购买方名称",
    "购买方税号",
    "销售方名称",
    "销售方税号",
    "未税金额",
    "税额",
    "价税合计",
    "币种",
    "校验状态",
    "备注",
]
PROMPT = """识别中国增值税发票。只输出 JSON，不要 Markdown。格式：
{"invoices":[{"invoice_type":"电子发票","invoice_number":"","issue_date":"YYYY-MM-DD","buyer_name":"","buyer_tax_id":"","seller_name":"","seller_tax_id":"","amount_excluding_tax":0.00,"tax_amount":0.00,"total_amount":0.00,"currency":"CNY","notes":""}]}
数字字段必须是数字。看不清的文本填空字符串，不能猜测；一张图没有发票则 invoices 为空数组。"""


def emit(
    kind: str,
    payload: dict[str, Any],
    seq: int,
    *,
    correlation: str = "",
    run: dict[str, Any] | None = None,
) -> None:
    """Write one UTF-8 NDJSON LAP frame to standard output."""
    frame: dict[str, Any] = {
        "lap": "0.1",
        "id": f"{AGENT_ID}:{seq}",
        "producer": AGENT_ID,
        "seq": seq,
        "type": kind,
        "payload": payload,
    }
    if correlation:
        frame["correlation_id"] = correlation
    if run:
        frame["run"] = run
    print(json.dumps(frame, ensure_ascii=False, separators=(",", ":")), flush=True)


def fail(message: str, code: str = "LAP-201") -> dict[str, Any]:
    """Build one protocol-valid failed terminal result."""
    return {
        "status": "failed",
        "summary": "发票处理失败。",
        "error": {"code": code, "message": message, "retryable": code == "LAP-500"},
    }


def input_files(context: Any, run_root: Path) -> list[tuple[Path, str]]:
    """Resolve and verify Host-staged invoice input artifacts."""
    artifacts = context.get("artifacts") if isinstance(context, dict) else None
    if not isinstance(artifacts, list) or not artifacts:
        raise ValueError("请上传至少一个发票图片或 PDF。")
    if len(artifacts) > MAX_FILES:
        raise ValueError(f"一次最多处理 {MAX_FILES} 个附件。")
    resolved: list[tuple[Path, str]] = []
    for item in artifacts:
        if not isinstance(item, dict) or not isinstance(item.get("name"), str):
            raise ValueError("附件描述无效。")
        name = item["name"]
        if Path(name).name != name or Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
            raise ValueError(f"不支持的附件类型: {name}")
        match = INPUT_URI.fullmatch(str(item.get("uri", "")))
        if not match:
            raise ValueError(f"附件引用无效: {name}")
        path = run_root / "input" / match.group(1)
        if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
            raise ValueError(f"附件不存在或超过 10 MiB: {name}")
        expected = str(item.get("sha256", "")).lower()
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if expected and expected != actual:
            raise ValueError(f"附件摘要不匹配: {name}")
        resolved.append((path, name))
    return resolved


def image_pages(path: Path, name: str) -> list[tuple[bytes, str, int, str]]:
    """Read one image or render the bounded page set of one PDF."""
    suffix = path.suffix.lower()
    if suffix != ".pdf":
        media = mimetypes.types_map.get(suffix, "image/jpeg")
        return [(path.read_bytes(), media, 1, "")]
    pages: list[tuple[bytes, str, int, str]] = []
    with pymupdf.open(path) as document:
        if document.page_count > MAX_PDF_PAGES:
            raise ValueError(f"PDF 超过 {MAX_PDF_PAGES} 页: {name}")
        for index, page in enumerate(document, 1):
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            pages.append((pixmap.tobytes("png"), "image/png", index, page.get_text()))
    return pages


def vision(
    image: bytes, media_type: str, detail: str, text_layer: str = ""
) -> list[dict[str, Any]]:
    """Call the authorized vision endpoint and return invoice JSON objects."""
    key = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if not key:
        raise RuntimeError("缺少 DEEPSEEK_API_KEY。")
    url = os.environ.get(
        "DEEPSEEK_VISION_BASE_URL", "https://api.deepseek.com/chat/completions"
    )
    model = os.environ.get("DEEPSEEK_VISION_MODEL", "deepseek-flash")
    data_url = f"data:{media_type};base64,{base64.b64encode(image).decode('ascii')}"
    body = json.dumps(
        {
            "model": model,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": PROMPT
                            + (
                                "\nPDF 文本层如下，请与图像交叉核对，优先使用其中完整号码和名称：\n"
                                + text_layer
                                if text_layer
                                else ""
                            ),
                        },
                        {
                            "type": "image_url",
                            "image_url": {"url": data_url, "detail": detail},
                        },
                    ],
                }
            ],
            "response_format": {"type": "json_object"},
            "max_tokens": 4096,
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            payload = json.load(response)
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"DeepSeek 视觉请求失败: {exc}") from exc
    try:
        content = payload["choices"][0]["message"]["content"]
        invoices = json.loads(content)["invoices"]
    except (KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
        raise RuntimeError("DeepSeek 返回了无效的发票 JSON。") from exc
    if not isinstance(invoices, list):
        raise RuntimeError("DeepSeek 发票结果必须是数组。")
    return [item for item in invoices if isinstance(item, dict)]


def money(value: Any) -> Decimal | None:
    """Convert one model-provided monetary field to two decimal places."""
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def normalize(
    raw: dict[str, Any], source: str, page: int
) -> tuple[list[Any], list[str]]:
    """Validate one vision response and build a workbook row with warnings."""
    warnings: list[str] = []
    subtotal = money(raw.get("amount_excluding_tax"))
    tax = money(raw.get("tax_amount"))
    total = money(raw.get("total_amount"))
    if subtotal is None or tax is None or total is None:
        warnings.append("金额字段缺失或格式错误")
    elif abs((subtotal + tax) - total) > Decimal("0.02"):
        warnings.append("价税合计与金额加税额不一致")
    number = str(raw.get("invoice_number", "")).strip()
    if not number:
        warnings.append("缺少发票号码")
    elif not re.fullmatch(r"(?:\d{8}|\d{10}|\d{12}|\d{20})", number):
        warnings.append("发票号码长度或格式异常")
    row = [
        source,
        page,
        str(raw.get("invoice_type", "")).strip(),
        number,
        str(raw.get("issue_date", "")).strip(),
        str(raw.get("buyer_name", "")).strip(),
        str(raw.get("buyer_tax_id", "")).strip(),
        str(raw.get("seller_name", "")).strip(),
        str(raw.get("seller_tax_id", "")).strip(),
        float(subtotal) if subtotal is not None else None,
        float(tax) if tax is not None else None,
        float(total) if total is not None else None,
        str(raw.get("currency", "CNY") or "CNY"),
        "通过" if not warnings else "需复核",
        "; ".join(
            warnings + ([str(raw.get("notes", "")).strip()] if raw.get("notes") else [])
        ),
    ]
    return row, warnings


def workbook(rows: list[list[Any]], target: Path) -> None:
    """Write the validated invoice rows as a styled XLSX workbook."""
    book = Workbook()
    sheet = book.active
    sheet.title = "发票明细"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = (
            fill,
            Font(color="FFFFFF", bold=True),
            Alignment(horizontal="center"),
        )
    for column in (10, 11, 12):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "#,##0.00"
    widths = [24, 8, 18, 24, 14, 30, 24, 30, 24, 14, 12, 14, 10, 12, 36]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    book.close()


def execute(payload: dict[str, Any], run: dict[str, Any], seq: int) -> int:
    """Execute one accepted LAP invoice-extraction Run."""
    request_input = payload.get("input") or {}
    detail = (
        request_input.get("detail", "original")
        if isinstance(request_input, dict)
        else "original"
    )
    if detail not in {"original", "high", "auto", "low"}:
        emit("run.result", fail("detail 参数无效。"), seq, run=run)
        return seq
    try:
        root = Path(os.environ["LAP_RUN_ROOT"])
        sources = input_files(payload.get("context"), root)
        rows: list[list[Any]] = []
        warnings: list[str] = []
        for path, name in sources:
            for image, media, page, text_layer in image_pages(path, name):
                seq += 1
                emit(
                    "run.progress",
                    {"phase": "recognize", "message": f"正在识别 {name} 第 {page} 页"},
                    seq,
                    run=run,
                )
                for item in vision(image, media, detail, text_layer):
                    row, issues = normalize(item, name, page)
                    rows.append(row)
                    warnings.extend(f"{name} 第 {page} 页: {issue}" for issue in issues)
        if not rows:
            emit("run.result", fail("附件中未识别到发票。"), seq + 1, run=run)
            return seq + 1
        target = root / "output" / "vat-invoices.xlsx"
        workbook(rows, target)
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        artifact = {
            "id": "vat-invoices",
            "name": target.name,
            "media_type": XLSX_MEDIA,
            "uri": "lap://run/output/vat-invoices.xlsx",
            "size_bytes": target.stat().st_size,
            "sha256": digest,
            "delivery": {
                "kind": "download",
                "required": True,
                "disposition": "attachment",
                "scope": "session",
                "semantic": {
                    "role": "final_result",
                    "title": "增值税发票整理结果",
                    "description": f"已识别 {len(rows)} 张发票并整理为 Excel。",
                    "source_count": len(sources),
                },
            },
        }
        seq += 1
        emit("run.artifact", artifact, seq, run=run)
        seq += 1
        emit(
            "run.result",
            {
                "status": "succeeded",
                "summary": f"已整理 {len(rows)} 张发票。",
                "output": {
                    "invoice_count": len(rows),
                    "source_count": len(sources),
                    "workbook": target.name,
                    "warnings": warnings,
                },
                "artifacts": [artifact],
            },
            seq,
            run=run,
        )
    except ValueError as exc:
        seq += 1
        emit("run.result", fail(str(exc)), seq, run=run)
    except Exception as exc:
        print(f"invoice agent error: {exc}", file=sys.stderr, flush=True)
        seq += 1
        emit(
            "run.result",
            fail("发票识别服务执行失败，请检查模型配置或稍后重试。", "LAP-500"),
            seq,
            run=run,
        )
    return seq


def main() -> int:
    """Serve LAP Local frames from standard input until shutdown."""
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if callable(reconfigure):
            with contextlib.suppress(ValueError, OSError):
                reconfigure(encoding="utf-8")
    seq = 0
    for raw in sys.stdin:
        message = json.loads(raw)
        kind = message.get("type")
        if kind == "agent.hello":
            seq += 1
            emit(
                "agent.welcome",
                {
                    "selected_lap": "0.1",
                    "profiles": ["lap-local/0.1"],
                    "agent_id": AGENT_ID,
                    "version": VERSION,
                    "max_concurrency": 1,
                },
                seq,
                correlation=str(message.get("id", "")),
            )
        elif kind == "run.start":
            run = message.get("run") or {}
            seq += 1
            emit(
                "run.accepted",
                {"capability": "invoice.extract"},
                seq,
                correlation=str(message.get("id", "")),
                run=run,
            )
            seq = execute(message.get("payload") or {}, run, seq)
        elif kind == "agent.shutdown":
            return 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
